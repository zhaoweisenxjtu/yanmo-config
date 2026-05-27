#!/usr/bin/env python3
"""Add brand icons to the AI Agent training PPT"""
import cairosvg
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from io import BytesIO

PPT_PATH = "/root/.openclaw/workspace/deliverables/AI Agent深度培训手册_PPT.pptx"
ICON_DIR = "/tmp/ppt_icons"
os.makedirs(ICON_DIR, exist_ok=True)

# Brand icons: (name, bg_color, text_color, label)
BRANDS = {
    # Models
    "GPT-5.5": ("#10a37f", "#ffffff", "G"),
    "GPT": ("#10a37f", "#ffffff", "G"),
    "OpenAI": ("#10a37f", "#ffffff", "O"),
    "Claude Opus": ("#d97706", "#ffffff", "C"),
    "Claude": ("#d97706", "#ffffff", "C"),
    "Anthropic": ("#d97706", "#ffffff", "A"),
    "Gemini": ("#4285f4", "#ffffff", "G"),
    "Google": ("#4285f4", "#ffffff", "G"),
    "DeepSeek V4": ("#4f46e5", "#ffffff", "D"),
    "DeepSeek": ("#4f46e5", "#ffffff", "D"),
    "Grok": ("#000000", "#ffffff", "G"),
    "xAI": ("#000000", "#ffffff", "x"),
    "Llama": ("#0466c8", "#ffffff", "L"),
    "Meta": ("#0466c8", "#ffffff", "M"),
    "Qwen": ("#ff6a00", "#ffffff", "Q"),
    "阿里巴巴": ("#ff6a00", "#ffffff", "A"),
    "GLM-5": ("#8b5cf6", "#ffffff", "G"),
    "智谱AI": ("#8b5cf6", "#ffffff", "Z"),
    "Kimi": ("#f97316", "#ffffff", "K"),
    "月之暗面": ("#f97316", "#ffffff", "M"),
    "ERNIE": ("#2563eb", "#ffffff", "E"),
    "百度": ("#2563eb", "#ffffff", "B"),
    "MiniMax": ("#06b6d4", "#ffffff", "M"),
    "Baichuan": ("#0891b2", "#ffffff", "B"),
    "豆包": ("#e11d48", "#ffffff", "D"),
    "字节跳动": ("#e11d48", "#ffffff", "B"),
    "讯飞星火": ("#059669", "#ffffff", "X"),
    "科大讯飞": ("#059669", "#ffffff", "K"),
    # Workbenches
    "OpenClaw": ("#e74c3c", "#ffffff", "OC"),
    "ChatGPT Codex": ("#10a37f", "#ffffff", "CC"),
    "Claude Code": ("#d97706", "#ffffff", "CL"),
    "AutoGPT": ("#7c3aed", "#ffffff", "AG"),
    "Dify": ("#1677ff", "#ffffff", "D"),
    "Coze": ("#e11d48", "#ffffff", "C"),
    "通义灵码": ("#ff6a00", "#ffffff", "TL"),
    "百度超级助理": ("#2563eb", "#ffffff", "BS"),
    "腾讯 WorkBuddy": ("#06b6d4", "#ffffff", "TW"),
    "LangChain": ("#059669", "#ffffff", "LC"),
    "CrewAI": ("#f97316", "#ffffff", "CA"),
    "MetaGPT": ("#8b5cf6", "#ffffff", "MG"),
    "n8n": ("#e11d48", "#ffffff", "N"),
    "Flowise": ("#0891b2", "#ffffff", "F"),
    "Cursor": ("#6d28d9", "#ffffff", "C"),
    "Windsurf": ("#0d9488", "#ffffff", "W"),
}

def create_icon_svg(bg_color, text_color, label, size=32):
    """Create a simple circle-with-initial icon"""
    r = size // 2
    cx, cy = r, r
    font_size = size * 0.5
    return f'''<svg xmlns="http://www.w3.org/2000/svg" width="{size}" height="{size}" viewBox="0 0 {size} {size}">
  <circle cx="{cx}" cy="{cy}" r="{r-1}" fill="{bg_color}"/>
  <text x="{cx}" y="{cy+font_size*0.35}" text-anchor="middle" fill="{text_color}" 
        font-family="Arial,sans-serif" font-size="{font_size}" font-weight="bold">{label}</text>
</svg>'''

# Generate all icons
print("Generating brand icons...")
for name, (bg, text, label) in BRANDS.items():
    svg = create_icon_svg(bg, text, label)
    png_path = os.path.join(ICON_DIR, f"{name.replace(' ','_').replace('/','_')}.png")
    try:
        cairosvg.svg2png(bytestring=svg.encode(), write_to=png_path, output_width=64, output_height=64)
    except Exception as e:
        print(f"  ❌ {name}: {e}")
print(f"  Generated {len(BRANDS)} icons")

# Open the PPT and add icons to relevant slides
prs = Presentation(PPT_PATH)

def add_icon_to_shape(slide, icon_name, left, top, width=Inches(0.35), height=Inches(0.35)):
    """Add a brand icon PNG to a slide"""
    name_key = icon_name.replace(' ', '_').replace('/', '_').replace('.', '_')
    # Try name variants
    for key in [name_key, icon_name, icon_name.split()[0] if ' ' in icon_name else '']:
        fname = key.replace(' ', '_').replace('/', '_').replace('.', '_') + '.png'
        path = os.path.join(ICON_DIR, fname)
        if os.path.exists(path):
            pic = slide.shapes.add_picture(path, left, top, width, height)
            return pic
    return None

# Map: slide number -> list of (brand_name, position)
icon_map = {
    2: [],  # chapter page
    3: [("OpenAI", Inches(3.0), Inches(1.95)), ("AutoGPT", Inches(7.5), Inches(4.25))],
    4: [],  # workflows
    5: [],  # memory
    6: [],  # RAG
    7: [],  # skills/MCP
    8: [("CrewAI", Inches(4.0), Inches(4.9)), ("MetaGPT", Inches(8.0), Inches(4.9))],
    9: [],  # chapter
    10: [("GPT-5.5", Inches(1.7), Inches(1.55)), ("Claude Opus", Inches(4.5), Inches(1.55)),
         ("Gemini", Inches(7.2), Inches(1.55)), ("DeepSeek V4", Inches(9.9), Inches(1.55)),
         ("Grok", Inches(1.7), Inches(3.5)), ("Llama", Inches(4.5), Inches(3.5)),
         ("Qwen", Inches(7.2), Inches(3.5)), ("GLM-5", Inches(9.9), Inches(3.5))],
    11: [("DeepSeek V4", Inches(0.7), Inches(1.65)), ("Qwen", Inches(3.8), Inches(1.65)),
         ("GLM-5", Inches(6.8), Inches(1.65)), ("Kimi", Inches(9.8), Inches(1.65))],
    12: [],  # selection guide
    13: [],  # chapter
    14: [("OpenClaw", Inches(0.45), Inches(1.65)), ("ChatGPT Codex", Inches(3.0), Inches(1.65)),
         ("Claude Code", Inches(5.5), Inches(1.65)), ("AutoGPT", Inches(8.0), Inches(1.65)),
         ("Dify", Inches(10.5), Inches(1.65))],
    15: [("Dify", Inches(0.45), Inches(1.65)), ("Coze", Inches(3.0), Inches(1.65)),
         ("通义灵码", Inches(5.5), Inches(1.65)), ("百度超级助理", Inches(8.0), Inches(1.65)),
         ("腾讯 WorkBuddy", Inches(10.5), Inches(1.65))],
    16: [],  # comparison
    17: [],  # chapter
    18: [("OpenClaw", Inches(1.2), Inches(1.55))],
    19: [],  # scenarios
}

print("\nAdding icons to slides...")
icon_count = 0
for slide_num, icons in icon_map.items():
    if slide_num >= len(prs.slides):
        continue
    slide = prs.slides[slide_num]
    for brand, left, top in icons:
        result = add_icon_to_shape(slide, brand, left, top)
        if result:
            icon_count += 1
            print(f"  Slide {slide_num+1}: +{brand} icon at ({left}, {top})")
        else:
            print(f"  Slide {slide_num+1}: ❌ {brand} icon not found")

# Save updated PPT
output_path = "/root/.openclaw/workspace/deliverables/AI Agent深度培训手册_PPT.pptx"
prs.save(output_path)
print(f"\n✅ Updated PPT saved: {output_path}")
print(f"   Total icons added: {icon_count}")
print(f"   Total slides: {len(prs.slides)}")

import os
size = os.path.getsize(output_path)
print(f"   Size: {size/1024:.1f} KB")
