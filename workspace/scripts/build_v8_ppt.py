#!/usr/bin/env python3
"""
v8: Scale all fonts to 24/18/16/12pt hierarchy.
Also: scale text boxes, add brand icons (small), fix theme fonts, fill empty slides.
"""
import os, shutil, zipfile, re
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from cairosvg import svg2png

V1 = 'deliverables/AI Agent深度培训手册_清新科技风_20页.pptx'
SVG_SRC = 'deliverables/AI Agent深度培训手册_清新科技风_20页_svg版.pptx'
DST = 'deliverables/AI Agent深度培训手册_清新科技风_20页_v8.pptx'
TMP = DST.replace('.pptx', '_tmp.pptx')
BP = 'temp/brand_pngs'
SW, SH = 12192000, 6858000

def map_font(pt):
    """Map current font size to target hierarchy"""
    if pt >= 22: return 24
    elif pt >= 16: return 18
    elif pt >= 11: return 16
    elif pt >= 8: return 12
    else: return 10  # extra-small stays 10

def add_pic(slide, name, x, y, w, h):
    p = f'{BP}/{name}'
    if os.path.exists(p):
        try: slide.shapes.add_picture(p, x, y, w, h); return True
        except: pass
    return False

# ── Step 1: Fix theme font ──
print("🔤 Theme font...")
shutil.copy2(V1, TMP)
with zipfile.ZipFile(V1) as zin, zipfile.ZipFile(TMP,'w',zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        d = zin.read(item.filename)
        if item.filename == 'ppt/theme/theme1.xml':
            d = d.decode().replace('<a:ea typeface=""/>', '<a:ea typeface="Microsoft YaHei"/>').encode()
        zout.writestr(item, d)

prs = Presentation(TMP)

# ── Step 2: Scale fonts + text boxes ──
print("📏 Scaling fonts...")
for i, slide in enumerate(prs.slides):
    box_scales = []
    
    for s in slide.shapes:
        if not s.has_text_frame:
            continue
        
        tf = s.text_frame
        tf.word_wrap = True
        
        # Scale font sizes
        for p in tf.paragraphs:
            for r in p.runs:
                if r.font.size:
                    old = r.font.size.pt
                    new = map_font(old)
                    if new != old:
                        r.font.size = Pt(new)
                        box_scales.append(new / old)
            
            # Scale line spacing
            if p.line_spacing and hasattr(p.line_spacing, 'pt'):
                old_ls = p.line_spacing.pt
                new_ls = map_font(old_ls)
                if new_ls != old_ls:
                    p.line_spacing = Pt(new_ls)
        
        # Scale space before/after
        for p in tf.paragraphs:
            if p.space_before and p.space_before.pt:
                p.space_before = Pt(map_font(p.space_before.pt))
            if p.space_after and p.space_after.pt:
                p.space_after = Pt(map_font(p.space_after.pt))
    
        # Enlarge text box height based on font scale
        if box_scales and s.text_frame.text.strip():
            avg_s = sum(box_scales)/len(box_scales)
            new_h = min(int(s.height * avg_s * 1.15), SH - s.top - 10000)
            if new_h > s.height:
                s.height = new_h
    
    if box_scales:
        avg = sum(box_scales)/len(box_scales)
        print(f'  P{i+1}: avg_scale={avg:.2f}x')

# ── Step 3: Brand icons ──
brands_map = [
    # P08 评测排行
    (8, 'openai', 14.5, 21, 2.5, 2.5), (8, 'anthropic', 14.5, 27, 2.5, 2.5),
    (8, 'google', 14.5, 33, 2.5, 2.5), (8, 'deepseek', 14.5, 39, 2.5, 2.5),
    (8, 'meta', 14.5, 45, 2.5, 2.5), (8, 'xai', 14.5, 51, 2.5, 2.5),
    (8, 'alibaba', 14.5, 57, 2.5, 2.5), (8, 'cohere_badge', 14.5, 63, 2.5, 2.5),
    (8, 'ibm', 14.5, 69, 2.5, 2.5), (8, 'mistral', 49, 21, 2.5, 2.5),
    (8, 'alibaba', 49, 27, 2.5, 2.5), (8, 'baidu', 49, 33, 2.5, 2.5),
    (8, 'bytedance', 49, 39, 2.5, 2.5), (8, 'tencent', 49, 45, 2.5, 2.5),
    (8, 'kimi_badge', 49, 51, 2.5, 2.5), (8, 'yi_badge', 49, 57, 2.5, 2.5),
    # P09 第一梯队 SVG overlay
    (9, 'openai', 7, 22, 3, 3), (9, 'anthropic', 7, 42, 3, 3),
    (9, 'google', 7, 62, 3, 3), (9, 'deepseek', 72, 22, 3, 3),
    # P10 中国模型
    (10, 'deepseek', 5, 38, 2.5, 2.5), (10, 'alibaba', 5, 50, 2.5, 2.5),
    (10, 'baidu', 5, 62, 2.5, 2.5), (10, 'bytedance', 50, 38, 2.5, 2.5),
    (10, 'tencent', 50, 50, 2.5, 2.5), (10, 'kimi_badge', 50, 62, 2.5, 2.5),
    # P11 全球工作台
    (11, 'github', 2.5, 42, 2.5, 2.5), (11, 'anthropic', 25.5, 42, 2.5, 2.5),
    (11, 'cursor', 48.5, 42, 2.5, 2.5), (11, 'dify', 71.5, 42, 2.5, 2.5),
    (11, 'coze', 87.5, 42, 2.5, 2.5),
    # P12 中国工作台
    (12, 'dify', 5.5, 42, 2.5, 2.5), (12, 'coze', 33.5, 42, 2.5, 2.5),
    (12, 'autogpt_badge', 59.5, 42, 2.5, 2.5),
    (12, 'lovable_badge', 19.5, 68, 2.5, 2.5), (12, 'bolt_badge', 45.5, 68, 2.5, 2.5),
    (12, 'v0_badge', 71.5, 68, 2.5, 2.5),
    # P13 OpenClaw
    (13, 'github', 82, 3, 2, 2),
]

print("\n🏷️ Brand icons...")
for sn, name, xp, yp, wp, hp in brands_map:
    sl = prs.slides[sn-1]
    x = int(SW * xp / 100); y = int(SH * yp / 100)
    w = int(SW * wp / 100); h = int(SH * hp / 100)
    if not add_pic(sl, f'{name}.png', x, y, w, h):
        pass  # silently skip missing

# ── Step 4: Fill empty slides 6,9 with scaled SVG content ──
print("🖼️ SVG slides...")
with zipfile.ZipFile(SVG_SRC) as z:
    for sn in [6, 9]:
        svg = z.read(f'ppt/media/image{sn}.svg').decode()
        # Scale font sizes in SVG: multiply all font-size values by ~1.5x
        def scale_svg_font(m):
            val = float(m.group(1))
            new_val = val * 1.5
            return f'font-size="{new_val:.0f}"'
        svg = re.sub(r'font-size="([\d.]+)"', scale_svg_font, svg)
        
        # Fix font-family
        svg = svg.replace(
            "font-family=\"'Noto Sans SC','Microsoft YaHei',sans-serif\"",
            "font-family=\"'Noto Sans CJK SC','WenQuanYi Micro Hei',sans-serif\""
        )
        svg = re.sub(r'(\w+)="([^"]*)"(\s+\1="[^"]*")+', lambda m: f'{m.group(1)}="{m.group(2)}"', svg)
        
        png = svg2png(bytestring=svg.encode(), scale=2, output_width=2560, output_height=1440)
        pp = f'/tmp/s{sn}.png'
        open(pp, 'wb').write(png)
        sl = prs.slides[sn-1]
        pic = sl.shapes.add_picture(pp, 0, 0, SW, SH)
        spTree = sl.shapes._spTree
        spTree.remove(pic._element)
        spTree.insert(2, pic._element)
        print(f'  P{sn}: {len(png)}B')

# ── Step 5: Save ──
prs.save(DST)
os.remove(TMP)

# Verify
p = Presentation(DST)
pics = sum(1 for s in p.slides for sh in s.shapes if sh.shape_type == 13)
empty = sum(1 for s in p.slides if len(list(s.shapes)) <= 1)
print(f"\n📊 Slides: {len(p.slides)}, Pics: {pics}, Near-empty: {empty}")
print(f"📦 Size: {os.path.getsize(DST):,}B")
print(f"✅ Done: {DST}")
